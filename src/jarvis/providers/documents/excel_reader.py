from __future__ import annotations

import csv
import logging
from io import BytesIO, StringIO
from typing import Any

logger = logging.getLogger(__name__)

_MAX_ROWS_PER_SHEET = 200
_MAX_CELL_LEN = 120


def read_excel(data: bytes, *, file_name: str = "workbook.xlsx") -> tuple[str, dict[str, Any], list[str]]:
    """Lit toutes les feuilles Excel et produit un résumé + échantillon structuré."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl requis pour lire les fichiers Excel.") from exc

    warnings: list[str] = []
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    sheet_names = list(wb.sheetnames)
    metadata: dict[str, Any] = {
        "sheets": sheet_names,
        "sheet_summaries": [],
        "truncated": False,
    }
    parts: list[str] = [f"Fichier Excel : {file_name}", f"Feuilles : {', '.join(sheet_names)}"]

    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        headers = [_cell_str(c) for c in (header_row or [])]
        headers = [h for h in headers if h]

        row_count = 0
        sample_rows: list[list[str]] = []
        for row in rows_iter:
            row_count += 1
            if len(sample_rows) < _MAX_ROWS_PER_SHEET:
                sample_rows.append([_cell_str(c) for c in row])

        summary = {
            "name": sheet_name,
            "columns": headers,
            "row_count": row_count,
            "sample_rows": len(sample_rows),
        }
        metadata["sheet_summaries"].append(summary)

        block = [f"\n## Feuille: {sheet_name}"]
        if headers:
            block.append(f"Colonnes ({len(headers)}) : {', '.join(headers[:30])}")
        block.append(f"Lignes de données (hors en-tête) : {row_count}")

        if row_count > _MAX_ROWS_PER_SHEET:
            metadata["truncated"] = True
            warnings.append(
                f"Feuille « {sheet_name} » tronquée à {_MAX_ROWS_PER_SHEET} lignes pour le contexte LLM."
            )

        if headers:
            block.append(" | ".join(headers))
        for row in sample_rows:
            if any(cell for cell in row):
                block.append(" | ".join(row[: len(headers) or len(row)]))

        parts.append("\n".join(block))

    wb.close()
    full = "\n".join(parts).strip()
    logger.info("Excel extrait file=%s sheets=%s chars=%s", file_name, len(sheet_names), len(full))
    return full, metadata, warnings


def read_csv(data: bytes, *, file_name: str = "data.csv") -> tuple[str, dict[str, Any], list[str]]:
    """Convertit un CSV en tableau texte avec en-têtes."""
    warnings: list[str] = []
    text = _decode_text_bytes(data)
    delimiter = _guess_delimiter(text)
    reader = csv.reader(StringIO(text), delimiter=delimiter)
    rows = list(reader)
    if not rows:
        return "", {"rows": 0, "columns": []}, ["Fichier CSV vide."]

    headers = [_cell_str(c) for c in rows[0]]
    data_rows = rows[1:]
    truncated = len(data_rows) > _MAX_ROWS_PER_SHEET
    if truncated:
        data_rows = data_rows[:_MAX_ROWS_PER_SHEET]
        warnings.append(f"CSV tronqué à {_MAX_ROWS_PER_SHEET} lignes.")

    parts = [
        f"Fichier CSV : {file_name}",
        f"Colonnes : {', '.join(headers)}",
        f"Lignes : {len(rows) - 1}",
        " | ".join(headers),
    ]
    for row in data_rows:
        parts.append(" | ".join(_cell_str(c) for c in row))

    metadata = {"rows": len(rows) - 1, "columns": headers, "delimiter": delimiter}
    return "\n".join(parts), metadata, warnings


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    s = str(value).strip().replace("\n", " ")
    return s[:_MAX_CELL_LEN]


def _decode_text_bytes(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _guess_delimiter(text: str) -> str:
    first_line = text.splitlines()[0] if text else ""
    if first_line.count(";") > first_line.count(","):
        return ";"
    if "\t" in first_line:
        return "\t"
    return ","
