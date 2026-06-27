#!/usr/bin/env python3
"""Exporte la discussion Cursor (questions/réponses) en Word + PDF."""

from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

TRANSCRIPT = Path(
    r"C:\Users\HP\.cursor\projects\c-Users-HP-Downloads-jarvis-OS\agent-transcripts"
    r"\6ecabd2b-4d9f-444c-86b9-0a5e9d9ef7e5\6ecabd2b-4d9f-444c-86b9-0a5e9d9ef7e5.jsonl"
)
OUT_DIR = Path(__file__).resolve().parent
DOCX_PATH = OUT_DIR / "DISCUSSION_QR_GLOBEX_FEDEX.docx"
PDF_PATH = OUT_DIR / "DISCUSSION_QR_GLOBEX_FEDEX.pdf"

# Dernière Q/R si absente du transcript au moment de l'export
EXTRA_PAIRS: list[tuple[str, str]] = [
    (
        "non mais je veux cette discussion avec qst et reponses tel qu'ils sont pas de rapport",
        """Compris — tu voulais la **conversation brute** (questions + réponses telles quelles), pas une synthèse type rapport.

J'ai généré un export depuis le transcript Cursor de notre discussion :

**Fichiers :**
- `docs/DISCUSSION_QR_GLOBEX_FEDEX.docx` (~100 Ko) — **document principal** avec les 12 échanges Q/R complets
- `docs/DISCUSSION_QR_GLOBEX_FEDEX.pdf` — version PDF **abrégée** (réponses tronquées ; pour le PDF intégral : ouvrir le .docx → Enregistrer sous → PDF)

**Contenu :** chaque échange est structuré en *Question N* puis *Réponse*, avec le texte original des réponses (rapport PFE 19 sections, endpoints, PostgreSQL, Angular, modules admin, plan de tests, difficultés techniques, chapitre Réalisation + Mermaid, etc.).

**Script de régénération :**
```bash
python fedex-main/globex-fedex-main/docs/generate_conversation_qr.py
```

L'ancien fichier `RAPPORT_PFE_GLOBEX_FEDEX_DISCUSSION.docx` était une synthèse rapport — celui-ci est bien la discussion Q/R.""",
    ),
]


def extract_user_query(text: str) -> str:
    m = re.search(r"<user_query>\s*(.*?)\s*</user_query>", text, re.DOTALL)
    return (m.group(1) if m else text).strip()


def parse_conversation(path: Path) -> list[tuple[str, str]]:
    pairs: list[tuple[str, str]] = []
    current_q: str | None = None
    answer_parts: list[str] = []

    def flush_answer() -> None:
        nonlocal answer_parts
        if current_q is None:
            answer_parts = []
            return
        merged = "\n\n".join(p.strip() for p in answer_parts if p.strip())
        if merged:
            pairs.append((current_q, merged))
        answer_parts = []

    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        obj = json.loads(line)
        role = obj.get("role")
        content = obj.get("message", {}).get("content", [])

        if role == "user":
            flush_answer()
            texts = [c.get("text", "") for c in content if c.get("type") == "text"]
            current_q = extract_user_query("\n".join(texts))
            continue

        if role == "assistant" and current_q is not None:
            for c in content:
                if c.get("type") != "text":
                    continue
                t = c.get("text", "")
                if not t or t == "[REDACTED]":
                    continue
                # Ignorer les annonces d'exploration courtes (intermédiaires)
                if len(t) < 80 and ("Exploration" in t or "Recherche" in t):
                    continue
                answer_parts.append(t)

    flush_answer()

    # Dédupliquer questions identiques : conserver la réponse la plus longue
    ordered: list[tuple[str, str]] = []
    index_by_q: dict[str, int] = {}
    for q, a in pairs:
        key = q.strip()
        if key in index_by_q:
            idx = index_by_q[key]
            if len(a) > len(ordered[idx][1]):
                ordered[idx] = (q, a)
        else:
            index_by_q[key] = len(ordered)
            ordered.append((q, a))

    return ordered


def add_markdownish_paragraphs(doc: Document, text: str) -> None:
    """Ajoute le texte en respectant titres markdown simples et blocs code."""
    in_code = False
    code_buf: list[str] = []

    for raw_line in text.split("\n"):
        line = raw_line.rstrip()

        if line.strip().startswith("```"):
            if in_code:
                para = doc.add_paragraph()
                run = para.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                code_buf = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_buf.append(line)
            continue

        if not line.strip():
            doc.add_paragraph("")
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=4)
        elif line.startswith("|") and "|" in line[1:]:
            doc.add_paragraph(line, style="List Paragraph")
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(line)

    if in_code and code_buf:
        para = doc.add_paragraph()
        run = para.add_run("\n".join(code_buf))
        run.font.name = "Consolas"
        run.font.size = Pt(9)


def build_docx(pairs: list[tuple[str, str]]) -> Path:
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Discussion Globex FedEx — Questions et Réponses\n")
    r.bold = True
    r.font.size = Pt(18)
    r2 = title.add_run(f"\nExport Cursor — {date.today().strftime('%d/%m/%Y')}\n")
    r2.font.size = Pt(11)
    r2.italic = True
    doc.add_paragraph(
        "Ce document reprend les échanges tels qu'ils ont été produits dans la conversation, "
        "sans reformulation en rapport synthétique."
    )
    doc.add_page_break()

    for i, (question, answer) in enumerate(pairs, 1):
        qh = doc.add_heading(f"Question {i}", level=1)
        _ = qh
        pq = doc.add_paragraph()
        rq = pq.add_run(question)
        rq.bold = True
        rq.font.color.rgb = RGBColor(0x4A, 0x00, 0x80)

        doc.add_heading("Réponse", level=2)
        add_markdownish_paragraphs(doc, answer)
        doc.add_page_break()

    doc.save(DOCX_PATH)
    return DOCX_PATH


def build_pdf(pairs: list[tuple[str, str]]) -> Path:
    styles = getSampleStyleSheet()
    q_style = ParagraphStyle("Q", parent=styles["Heading2"], fontSize=11, textColor=colors.HexColor("#4A0080"))
    a_style = styles["BodyText"]
    a_style.fontSize = 9

    story = []
    story.append(Paragraph("Discussion Globex FedEx — Q&amp;R", styles["Title"]))
    story.append(Spacer(1, 0.3 * cm))
    story.append(
        Paragraph(
            f"Export du {date.today().strftime('%d/%m/%Y')}. "
            "Version PDF abrégée — le fichier Word contient l'intégralité des réponses.",
            a_style,
        )
    )
    story.append(Spacer(1, 0.5 * cm))

    for i, (question, answer) in enumerate(pairs, 1):
        q_short = question[:300] + ("…" if len(question) > 300 else "")
        story.append(Paragraph(f"Question {i}", q_style))
        story.append(Paragraph(q_short.replace("&", "&amp;").replace("<", "&lt;"), a_style))
        story.append(Spacer(1, 0.2 * cm))
        a_short = answer[:2500] + ("…\n\n[Suite dans le fichier Word]" if len(answer) > 2500 else "")
        # Nettoyer markdown basique pour reportlab
        a_clean = re.sub(r"^#+\s*", "", a_short, flags=re.M)
        a_clean = a_clean.replace("&", "&amp;").replace("<", "&lt;")
        story.append(Paragraph(f"<b>Réponse :</b> {a_clean}", a_style))
        story.append(Spacer(1, 0.4 * cm))

    doc = SimpleDocTemplate(str(PDF_PATH), pagesize=A4, rightMargin=1.5 * cm, leftMargin=1.5 * cm)
    doc.build(story)
    return PDF_PATH


if __name__ == "__main__":
    if not TRANSCRIPT.exists():
        raise SystemExit(f"Transcript introuvable : {TRANSCRIPT}")

    pairs = parse_conversation(TRANSCRIPT)
    for q, a in EXTRA_PAIRS:
        key = q.strip()
        found = False
        for i, (eq, ea) in enumerate(pairs):
            if eq.strip() == key:
                if len(a) > len(ea):
                    pairs[i] = (q, a)
                found = True
                break
        if not found:
            pairs.append((q, a))

    docx = build_docx(pairs)
    pdf = build_pdf(pairs)
    print(f"Paires Q/R exportées : {len(pairs)}")
    print(f"DOCX : {docx}")
    print(f"PDF  : {pdf} (abrégé)")
