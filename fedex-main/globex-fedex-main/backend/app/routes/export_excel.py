from io import BytesIO
import json

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.models.shipment_cache import ShipmentCache
from app.models.tracking_request import TrackingRequest
from app.models.user import User
from app.routes.deps import get_current_user
from app.services.activity_log_service import client_ip, write_log
from app.services.ai_assistant.export_dataset_cache import get_owner_export_dataset
from app.services.chat_export_service import TRACKING_ONLY_COLUMNS, generate_tracking_history_pdf_bytes, latest_tracking_rows
from app.services.quota_service import enforce_daily_quota
from app.services.user_notification_service import maybe_notify_export_ready

router = APIRouter(tags=["export"])

_HEADER_BY_LANG: dict[str, dict[str, str]] = {
    "fr": {
        "id": "ID",
        "tracking_number": "Numéro de suivi",
        "user_question": "Question utilisateur",
        "bot_response": "Réponse bot",
        "status": "Statut",
        "current_location": "Localisation",
        "estimated_delivery": "Livraison estimée",
        "created_at": "Créé le",
    },
    "en": {
        "id": "ID",
        "tracking_number": "Tracking number",
        "user_question": "User question",
        "bot_response": "Bot reply",
        "status": "Status",
        "current_location": "Location",
        "estimated_delivery": "Estimated delivery",
        "created_at": "Created at",
    },
    "ar": {
        "id": "المعرف",
        "tracking_number": "رقم التتبع",
        "user_question": "سؤال المستخدم",
        "bot_response": "رد المساعد",
        "status": "الحالة",
        "current_location": "الموقع",
        "estimated_delivery": "التسليم المتوقع",
        "created_at": "تاريخ الإنشاء",
    },
}
_ALL_COLUMNS: tuple[str, ...] = (
    "id",
    "tracking_number",
    "user_question",
    "bot_response",
    "status",
    "current_location",
    "estimated_delivery",
    "created_at",
)


def _normalize_lang(lang: str | None) -> str:
    raw = (lang or "fr").strip().lower()[:2]
    return raw if raw in _HEADER_BY_LANG else "fr"


def _normalize_columns(columns: str | None, preset: str | None = None) -> list[str]:
    if preset == "tracking":
        return list(TRACKING_ONLY_COLUMNS)
    if not columns:
        return list(_ALL_COLUMNS)
    parts = [p.strip().lower() for p in columns.split(",") if p.strip()]
    selected = [p for p in parts if p in _ALL_COLUMNS]
    return selected or list(_ALL_COLUMNS)


def _extract_events_from_cache(raw_json: str) -> list[dict]:
    try:
        data = json.loads(raw_json or "{}")
    except json.JSONDecodeError:
        return []
    events = data.get("events")
    if isinstance(events, list):
        return [ev for ev in events if isinstance(ev, dict)]
    return []


def _value_for_column(r: TrackingRequest, column: str):
    if column == "created_at":
        return r.created_at.isoformat() if r.created_at else ""
    return getattr(r, column, "")


@router.get("/export/tracking-history.xlsx")
def export_tracking_history(
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    limit: int = Query(default=100, ge=1, le=500),
    lang: str = Query(default="fr", max_length=8),
    columns: str | None = Query(default=None, max_length=400),
    preset: str | None = Query(default=None, max_length=32),
    session_id: int | None = Query(default=None, ge=1),
    tracking_numbers: str | None = Query(default=None, max_length=2000),
    include_events: bool = Query(default=False),
) -> StreamingResponse:
    enforce_daily_quota(db, user, kind="exports")
    lang_code = _normalize_lang(lang)
    selected_columns = _normalize_columns(columns, preset=preset)
    header_map = _HEADER_BY_LANG[lang_code]

    tn_list: list[str] | None = None
    if tracking_numbers:
        tn_list = [p.strip() for p in tracking_numbers.split(",") if p.strip()]

    if session_id is not None or tn_list:
        rows = latest_tracking_rows(
            db,
            user_id=user.id,
            session_id=session_id,
            tracking_numbers=tn_list,
            limit=limit,
        )
    else:
        stmt = (
            select(TrackingRequest)
            .where(TrackingRequest.user_id == user.id)
            .order_by(TrackingRequest.created_at.desc())
            .limit(limit)
        )
        rows = list(db.scalars(stmt).all())

    filename = "tracking-export.xlsx" if preset == "tracking" else "tracking-history.xlsx"

    write_log(
        db,
        action="export.tracking_excel",
        message=f"Export Excel historique suivi ({len(rows)} lignes)",
        category="system",
        level="INFO",
        user_id=user.id,
        ip_address=client_ip(request),
        metadata={
            "row_count": len(rows),
            "limit": limit,
            "lang": lang_code,
            "columns": ",".join(selected_columns),
            "include_events": include_events,
            "session_id": session_id,
            "tracking_numbers": tracking_numbers,
            "preset": preset,
        },
        commit=True,
    )
    maybe_notify_export_ready(db, user_id=user.id, row_count=len(rows))
    db.commit()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Historique"
    ws.append([header_map[c] for c in selected_columns])
    header_fill = PatternFill(fill_type="solid", fgColor="4D148C")
    for i, _ in enumerate(selected_columns, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in rows:
        ws.append([_value_for_column(r, col) for col in selected_columns])

    # Largeurs + wrap auto + auto-filter (tri/filtre natif Excel)
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    for i, col_name in enumerate(selected_columns, start=1):
        letter = get_column_letter(i)
        width = 20
        if col_name in {"user_question", "bot_response"}:
            width = 42
        elif col_name in {"tracking_number"}:
            width = 24
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Mise en forme conditionnelle simple sur la colonne statut
    if "status" in selected_columns and ws.max_row >= 2:
        status_col = selected_columns.index("status") + 1
        col_letter = get_column_letter(status_col)
        ref = f"{col_letter}2:{col_letter}{ws.max_row}"
        ws.conditional_formatting.add(
            ref,
            FormulaRule(
                formula=[f'ISNUMBER(SEARCH("livré",{col_letter}2))'],
                stopIfTrue=False,
                fill=PatternFill(fill_type="solid", fgColor="D1FAE5"),
            ),
        )
        ws.conditional_formatting.add(
            ref,
            FormulaRule(
                formula=[f'ISNUMBER(SEARCH("delivered",{col_letter}2))'],
                stopIfTrue=False,
                fill=PatternFill(fill_type="solid", fgColor="D1FAE5"),
            ),
        )
        ws.conditional_formatting.add(
            ref,
            FormulaRule(
                formula=[f'ISNUMBER(SEARCH("transit",{col_letter}2))'],
                stopIfTrue=False,
                fill=PatternFill(fill_type="solid", fgColor="FEF3C7"),
            ),
        )
        ws.conditional_formatting.add(
            ref,
            FormulaRule(
                formula=[f'ISNUMBER(SEARCH("exception",{col_letter}2))'],
                stopIfTrue=False,
                fill=PatternFill(fill_type="solid", fgColor="FEE2E2"),
            ),
        )

    if include_events:
        ws_events = wb.create_sheet(title="Events")
        event_headers = {
            "fr": ["Numéro de suivi", "Date", "Événement", "Lieu"],
            "en": ["Tracking number", "Date", "Event", "Location"],
            "ar": ["رقم التتبع", "التاريخ", "الحدث", "المكان"],
        }[lang_code]
        ws_events.append(event_headers)
        for i in range(1, 5):
            c = ws_events.cell(row=1, column=i)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
        cache_by_tn = {
            c.tracking_number: c
            for c in db.scalars(
                select(ShipmentCache).where(ShipmentCache.tracking_number.in_([r.tracking_number for r in rows]))
            ).all()
        }
        for r in rows:
            cache = cache_by_tn.get(r.tracking_number)
            events = _extract_events_from_cache(cache.raw_response_json if cache else "")
            if not events:
                ws_events.append([r.tracking_number, "", "", ""])
                continue
            for ev in events:
                ws_events.append(
                    [
                        r.tracking_number,
                        str(ev.get("at") or ""),
                        str(ev.get("description") or ""),
                        str(ev.get("location") or ""),
                    ]
                )
        ws_events.auto_filter.ref = ws_events.dimensions
        ws_events.freeze_panes = "A2"
        ws_events.column_dimensions["A"].width = 24
        ws_events.column_dimensions["B"].width = 24
        ws_events.column_dimensions["C"].width = 38
        ws_events.column_dimensions["D"].width = 26

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@router.get("/export/tracking-history.pdf")
def export_tracking_history_pdf(
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    limit: int = Query(default=100, ge=1, le=500),
    lang: str = Query(default="fr", max_length=8),
    preset: str | None = Query(default=None, max_length=32),
    session_id: int | None = Query(default=None, ge=1),
    tracking_numbers: str | None = Query(default=None, max_length=2000),
    include_events: bool = Query(default=True),
) -> StreamingResponse:
    enforce_daily_quota(db, user, kind="exports")
    lang_code = _normalize_lang(lang)

    tn_list: list[str] | None = None
    if tracking_numbers:
        tn_list = [p.strip() for p in tracking_numbers.split(",") if p.strip()]

    pdf_bytes, filename = generate_tracking_history_pdf_bytes(
        db,
        user_id=user.id,
        session_id=session_id,
        tracking_numbers=tn_list,
        limit=limit,
        preset=preset,
        include_events=include_events,
    )

    write_log(
        db,
        action="export.tracking_pdf",
        message=f"Export PDF historique suivi ({filename})",
        category="system",
        level="INFO",
        user_id=user.id,
        ip_address=client_ip(request),
        metadata={
            "limit": limit,
            "lang": lang_code,
            "include_events": include_events,
            "session_id": session_id,
            "tracking_numbers": tracking_numbers,
            "preset": preset,
            "filename": filename,
        },
        commit=True,
    )
    maybe_notify_export_ready(db, user_id=user.id, row_count=1)
    db.commit()

    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(
        BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers=headers,
    )


@router.get("/export/text.pdf")
def export_text_pdf(
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    export_token: str = Query(..., min_length=8, max_length=64),
) -> StreamingResponse:
    """Télécharge un PDF texte libre généré par l'assistant client (cache token)."""
    enforce_daily_quota(db, user, kind="exports")
    cached = get_owner_export_dataset(export_token, owner_id=user.id)
    if not cached:
        raise HTTPException(status_code=404, detail="Export introuvable ou expiré.")
    pdf_bytes = cached.get("pdf_bytes")
    if not isinstance(pdf_bytes, (bytes, bytearray)) or not pdf_bytes:
        raise HTTPException(status_code=404, detail="PDF introuvable.")
    filename = str(cached.get("filename") or "document.pdf")

    write_log(
        db,
        action="export.text_pdf",
        message=f"Export PDF texte client ({filename})",
        category="system",
        level="INFO",
        user_id=user.id,
        ip_address=client_ip(request),
        metadata={"export_token": export_token[:8], "filename": filename},
        commit=True,
    )
    maybe_notify_export_ready(db, user_id=user.id, row_count=1)
    db.commit()

    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(
        BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers=headers,
    )


@router.get("/export/text.xlsx")
def export_text_xlsx(
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    export_token: str = Query(..., min_length=8, max_length=64),
) -> StreamingResponse:
    """Télécharge un Excel généré par l'assistant client (cache token)."""
    enforce_daily_quota(db, user, kind="exports")
    cached = get_owner_export_dataset(export_token, owner_id=user.id)
    if not cached:
        raise HTTPException(status_code=404, detail="Export introuvable ou expiré.")
    xlsx_bytes = cached.get("xlsx_bytes")
    if not isinstance(xlsx_bytes, (bytes, bytearray)) or not xlsx_bytes:
        raise HTTPException(status_code=404, detail="Fichier Excel introuvable.")
    filename = str(cached.get("filename") or "export.xlsx")

    write_log(
        db,
        action="export.text_xlsx",
        message=f"Export Excel client ({filename})",
        category="system",
        level="INFO",
        user_id=user.id,
        ip_address=client_ip(request),
        metadata={"export_token": export_token[:8], "filename": filename},
        commit=True,
    )
    maybe_notify_export_ready(db, user_id=user.id, row_count=1)
    db.commit()

    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(
        BytesIO(xlsx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )
