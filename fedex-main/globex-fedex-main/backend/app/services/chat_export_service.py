"""Export Excel ciblé depuis une conversation chat."""

from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from io import BytesIO
from typing import Any

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.chat_message import ChatMessage, MessageSender
from app.models.shipment_cache import ShipmentCache
from app.models.tracking_request import TrackingRequest
from app.models.user import User
from app.services.chat_shipment_reply import is_session_follow_up
from app.utils.tracking_parser import is_plausible_tracking_number
from app.services.llm.tracking_extract import extract_all_tracking_numbers, extract_tracking_number
from app.services.message_attachment import unpack_message_text

TRACKING_ONLY_COLUMNS: tuple[str, ...] = (
    "tracking_number",
    "status",
    "current_location",
    "estimated_delivery",
    "created_at",
)

_EXPORT_TRIGGERS = (
    "export",
    "excel",
    "xlsx",
    "telecharger",
    "télécharger",
    "fichier",
    "historique",
    "génère",
    "genere",
    "générer",
    "generer",
)


def is_export_intent(text: str) -> bool:
    lowered = text.lower()
    if "pdf" in lowered:
        return True
    return any(word in lowered for word in _EXPORT_TRIGGERS) and (
        "excel" in lowered
        or "export" in lowered
        or "xlsx" in lowered
        or "fichier" in lowered
    )

_ALL_EXPORT_KEYWORDS = (
    "tous",
    "toutes",
    "all",
    "everything",
    "entier",
    "entière",
    "complete",
    "complet",
)


_TRACKING_MIN_DIGITS = 8


def _is_plausible_tracking_number(value: str) -> bool:
    return is_plausible_tracking_number(value)


def session_tracking_numbers(db: Session, session_id: int, user_id: int) -> list[str]:
    """Numéros uniques suivis dans la session (ordre chronologique d'apparition)."""
    from app.models.chat_message import ChatMessage
    from app.services.message_attachment import unpack_message_text

    rows = list(
        db.scalars(
            select(TrackingRequest)
            .where(
                TrackingRequest.session_id == session_id,
                TrackingRequest.user_id == user_id,
            )
            .order_by(TrackingRequest.created_at.asc())
        ).all()
    )
    seen: set[str] = set()
    ordered: list[str] = []
    for row in rows:
        if _is_plausible_tracking_number(row.tracking_number) and row.tracking_number not in seen:
            seen.add(row.tracking_number)
            ordered.append(row.tracking_number)

    if ordered:
        return ordered

    msg_rows = list(
        db.scalars(
            select(ChatMessage)
            .where(ChatMessage.session_id == session_id)
            .order_by(ChatMessage.created_at.asc())
        ).all()
    )
    for row in msg_rows:
        text, _, _ = unpack_message_text(row.message_text or "")
        for tn in extract_all_tracking_numbers(text or ""):
            if _is_plausible_tracking_number(tn) and tn not in seen:
                seen.add(tn)
                ordered.append(tn)
    return ordered


def user_recent_tracking_numbers(db: Session, user_id: int, limit: int = 5) -> list[str]:
    """N derniers numéros de suivi distincts du compte (ordre anti-chronologique)."""
    cap = max(1, min(limit, 50))
    rows = list(
        db.scalars(
            select(TrackingRequest.tracking_number)
            .where(TrackingRequest.user_id == user_id)
            .order_by(TrackingRequest.created_at.desc())
            .limit(cap * 4)
        ).all()
    )
    seen: set[str] = set()
    ordered: list[str] = []
    for tn in rows:
        if _is_plausible_tracking_number(tn) and tn not in seen:
            seen.add(tn)
            ordered.append(tn)
        if len(ordered) >= cap:
            break
    return ordered


_EXCEL_HEADER_BY_LANG: dict[str, dict[str, str]] = {
    "fr": {
        "tracking_number": "Numéro de suivi",
        "status": "Statut",
        "current_location": "Localisation",
        "estimated_delivery": "Livraison estimée",
        "created_at": "Créé le",
    },
    "en": {
        "tracking_number": "Tracking number",
        "status": "Status",
        "current_location": "Location",
        "estimated_delivery": "Estimated delivery",
        "created_at": "Created at",
    },
}


def _excel_value(row: TrackingRequest, column: str) -> str:
    if column == "created_at":
        return row.created_at.isoformat() if row.created_at else ""
    return str(getattr(row, column, "") or "")


def _extract_events_from_cache(raw_json: str) -> list[dict]:
    try:
        data = json.loads(raw_json or "{}")
    except json.JSONDecodeError:
        return []
    events = data.get("events")
    if isinstance(events, list):
        return [ev for ev in events if isinstance(ev, dict)]
    return []


def generate_tracking_excel_bytes(
    db: Session,
    rows: list[TrackingRequest],
    *,
    lang: str = "fr",
    include_events: bool = True,
) -> bytes:
    """Génère un fichier Excel en mémoire pour l'agent / e-mail."""
    lang_code = (lang or "fr").strip().lower()[:2]
    if lang_code not in _EXCEL_HEADER_BY_LANG:
        lang_code = "fr"
    header_map = _EXCEL_HEADER_BY_LANG[lang_code]
    columns = list(TRACKING_ONLY_COLUMNS)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Historique"
    ws.append([header_map[c] for c in columns])
    header_fill = PatternFill(fill_type="solid", fgColor="4D148C")
    for i in range(1, len(columns) + 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in rows:
        ws.append([_excel_value(r, col) for col in columns])

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    for i, col_name in enumerate(columns, start=1):
        letter = get_column_letter(i)
        width = 24 if col_name == "tracking_number" else 20
        ws.column_dimensions[letter].width = width

    if include_events and rows:
        ws_events = wb.create_sheet(title="Events")
        event_headers = {
            "fr": ["Numéro de suivi", "Date", "Événement", "Lieu"],
            "en": ["Tracking number", "Date", "Event", "Location"],
        }[lang_code]
        ws_events.append(event_headers)
        for i in range(1, 5):
            c = ws_events.cell(row=1, column=i)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = header_fill
        cache_by_tn = {
            c.tracking_number: c
            for c in db.scalars(
                select(ShipmentCache).where(
                    ShipmentCache.tracking_number.in_([r.tracking_number for r in rows])
                )
            ).all()
        }
        for r in rows:
            cache = cache_by_tn.get(r.tracking_number)
            events = _extract_events_from_cache(cache.raw_response_json if cache else "")
            if not events:
                ws_events.append([r.tracking_number, "", "", ""])
                continue
            for ev in events:
                ws_events.append(
                    [
                        r.tracking_number,
                        str(ev.get("at") or ""),
                        str(ev.get("description") or ""),
                        str(ev.get("location") or ""),
                    ]
                )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


_FEDEX_SUMMARY_HEADERS: dict[str, list[str]] = {
    "fr": [
        "Numéro de suivi",
        "Statut",
        "Localisation",
        "Livraison estimée",
        "Livraison effective",
        "Destinataire",
        "Origine",
        "Destination",
    ],
    "en": [
        "Tracking number",
        "Status",
        "Location",
        "Estimated delivery",
        "Actual delivery",
        "Recipient",
        "Origin",
        "Destination",
    ],
}

_FEDEX_EVENT_HEADERS: dict[str, list[str]] = {
    "fr": ["Date", "Événement", "Lieu"],
    "en": ["Date", "Event", "Location"],
}


def _shipment_from_fedex_payload(fedex_payload: dict[str, Any]) -> dict[str, Any]:
    if not fedex_payload.get("available"):
        return {
            k: fedex_payload[k]
            for k in ("tracking_number", "status", "current_location", "estimated_delivery", "message")
            if fedex_payload.get(k) is not None
        }
    shipment = fedex_payload.get("shipment")
    if isinstance(shipment, dict) and shipment:
        return shipment
    return {
        k: fedex_payload[k]
        for k in (
            "tracking_number",
            "status",
            "current_location",
            "estimated_delivery",
            "actual_delivery",
            "recipient",
            "origin_location",
            "destination_location",
            "events",
        )
        if fedex_payload.get(k) is not None
    }


def _apply_excel_header_style(ws, col_count: int) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="4D148C")
    for i in range(1, col_count + 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")


def generate_shipment_excel_from_fedex(
    fedex_payload: dict[str, Any],
    *,
    layout: str = "summary_and_events",
    lang: str = "fr",
) -> tuple[bytes, str]:
    """Génère un xlsx enrichi depuis les faits FedEx (sans DB)."""
    lang_code = (lang or "fr").strip().lower()[:2]
    if lang_code not in _FEDEX_SUMMARY_HEADERS:
        lang_code = "fr"
    shipment = _shipment_from_fedex_payload(fedex_payload)
    tn = str(shipment.get("tracking_number") or fedex_payload.get("tracking_number") or "colis")
    layout_norm = (layout or "summary_and_events").strip().lower()
    if layout_norm not in {"summary_only", "summary_and_events", "events_table"}:
        layout_norm = "summary_and_events"

    wb = openpyxl.Workbook()
    events = [ev for ev in (shipment.get("events") or []) if isinstance(ev, dict)][:40]

    if layout_norm == "events_table":
        ws = wb.active
        ws.title = "Chronologie" if lang_code == "fr" else "Timeline"
        headers = (
            ["Numéro de suivi"] + _FEDEX_EVENT_HEADERS[lang_code]
            if lang_code == "fr"
            else ["Tracking number"] + _FEDEX_EVENT_HEADERS[lang_code]
        )
        ws.append(headers)
        _apply_excel_header_style(ws, len(headers))
        if events:
            for ev in events:
                ws.append(
                    [
                        tn,
                        str(ev.get("at") or ev.get("occurred_at") or ""),
                        str(ev.get("description") or ""),
                        str(ev.get("location") or ""),
                    ]
                )
        else:
            msg = str(fedex_payload.get("message") or "Aucun scan disponible.")
            ws.append([tn, "", msg, ""])
    else:
        ws = wb.active
        ws.title = "Synthese" if lang_code == "fr" else "Summary"
        summary_headers = _FEDEX_SUMMARY_HEADERS[lang_code]
        ws.append(summary_headers)
        _apply_excel_header_style(ws, len(summary_headers))
        ws.append(
            [
                tn,
                str(shipment.get("status") or fedex_payload.get("status") or ""),
                str(shipment.get("current_location") or fedex_payload.get("current_location") or ""),
                str(shipment.get("estimated_delivery") or fedex_payload.get("estimated_delivery") or ""),
                str(shipment.get("actual_delivery") or ""),
                str(shipment.get("recipient") or ""),
                str(shipment.get("origin_location") or ""),
                str(shipment.get("destination_location") or ""),
            ]
        )
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        for i in range(1, len(summary_headers) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 22

        if layout_norm == "summary_and_events":
            ws_events = wb.create_sheet(title="Chronologie" if lang_code == "fr" else "Timeline")
            ev_headers = (
                ["Numéro de suivi"] + _FEDEX_EVENT_HEADERS[lang_code]
                if lang_code == "fr"
                else ["Tracking number"] + _FEDEX_EVENT_HEADERS[lang_code]
            )
            ws_events.append(ev_headers)
            _apply_excel_header_style(ws_events, len(ev_headers))
            if events:
                for ev in events:
                    ws_events.append(
                        [
                            tn,
                            str(ev.get("at") or ev.get("occurred_at") or ""),
                            str(ev.get("description") or ""),
                            str(ev.get("location") or ""),
                        ]
                    )
            else:
                ws_events.append([tn, "", str(fedex_payload.get("message") or ""), ""])
            ws_events.auto_filter.ref = ws_events.dimensions
            ws_events.freeze_panes = "A2"

    ts = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", tn[:30]).strip("-").lower() or "colis"
    filename = f"suivi-{slug}-{ts}.xlsx"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue(), filename


def latest_tracking_rows(
    db: Session,
    *,
    user_id: int,
    session_id: int | None = None,
    tracking_numbers: list[str] | None = None,
    limit: int = 100,
) -> list[TrackingRequest]:
    """Dernière entrée par numéro de suivi."""
    stmt = (
        select(TrackingRequest)
        .where(TrackingRequest.user_id == user_id)
        .order_by(TrackingRequest.created_at.desc())
    )
    if session_id is not None:
        stmt = stmt.where(TrackingRequest.session_id == session_id)
    if tracking_numbers:
        stmt = stmt.where(TrackingRequest.tracking_number.in_(tracking_numbers))
    stmt = stmt.limit(limit * 5)
    rows = list(db.scalars(stmt).all())
    by_tn: dict[str, TrackingRequest] = {}
    for row in rows:
        if row.tracking_number not in by_tn:
            by_tn[row.tracking_number] = row
    if tracking_numbers:
        return [by_tn[tn] for tn in tracking_numbers if tn in by_tn]
    return list(by_tn.values())[:limit]


def _last_bot_message(db: Session, session_id: int) -> ChatMessage | None:
    return db.scalars(
        select(ChatMessage)
        .where(
            ChatMessage.session_id == session_id,
            ChatMessage.sender == MessageSender.bot.value,
        )
        .order_by(ChatMessage.created_at.desc())
        .limit(1)
    ).first()


def _wants_all_trackings(message: str) -> bool:
    lowered = message.lower()
    return any(word in lowered for word in _ALL_EXPORT_KEYWORDS)


def parse_export_selection(message: str, trackings: list[str]) -> list[str] | None:
    """Interprète le choix utilisateur après une question de clarification."""
    if not trackings:
        return None
    if _wants_all_trackings(message):
        return list(trackings)

    for tn in extract_all_tracking_numbers(message):
        if tn in trackings:
            return [tn]

    stripped = message.strip()
    if re.fullmatch(r"\d+", stripped):
        idx = int(stripped) - 1
        if 0 <= idx < len(trackings):
            return [trackings[idx]]

    compact = re.sub(r"\s+", "", message)
    for tn in trackings:
        if tn in compact or tn in message:
            return [tn]

    return None


def _build_disambiguation_reply(trackings: list[str]) -> str:
    lines = [
        "Plusieurs colis ont été suivis dans cette conversation. Lequel voulez-vous exporter en Excel ?",
        "",
    ]
    for i, tn in enumerate(trackings, start=1):
        lines.append(f"{i}. **{tn}**")
    lines.extend(
        [
            "",
            "Répondez avec le **numéro** (1, 2, …), le **numéro de suivi**, ou tapez **tous** pour tout exporter dans un seul fichier.",
            "L'export contiendra uniquement les données de suivi (statut, localisation, ETA) — pas le texte de la conversation.",
        ]
    )
    return "\n".join(lines)


def resolve_export_targets(
    db: Session,
    session_id: int,
    user_id: int,
    message: str,
) -> tuple[list[str] | None, str | None, bool]:
    """
    Retourne (tracking_numbers, reply_if_prompt, should_download).
    - should_download True → lancer le téléchargement
    - reply_if_prompt → message bot (clarification ou erreur)
    - tracking_numbers → liste à exporter (vide si pas de download)
    """
    trackings = session_tracking_numbers(db, session_id, user_id)
    last_bot = _last_bot_message(db, session_id)

    if last_bot is not None and last_bot.source == "export_prompt":
        selected = parse_export_selection(message, trackings)
        if selected:
            return selected, None, True
        return (
            None,
            "Je n'ai pas compris votre choix. Indiquez le numéro dans la liste (ex. **1**), "
            "le numéro de suivi, ou **tous** pour exporter tous les colis de cette conversation.",
            False,
        )

    if not trackings:
        return (
            None,
            "Aucun colis n'a encore été suivi dans cette conversation. "
            "Demandez d'abord le suivi d'un numéro FedEx, puis demandez l'export Excel.",
            False,
        )

    if _wants_all_trackings(message):
        return list(trackings), None, True

    explicit = extract_tracking_number(message)
    if explicit and explicit in trackings:
        return [explicit], None, True

    found_in_msg = [tn for tn in extract_all_tracking_numbers(message) if tn in trackings]
    if len(found_in_msg) == 1:
        return [found_in_msg[0]], None, True
    if len(found_in_msg) > 1:
        return found_in_msg, None, True

    if len(trackings) == 1:
        return [trackings[0]], None, True

    return None, _build_disambiguation_reply(trackings), False


def handle_chat_export_request(
    db: Session,
    session,
    user: User,
    message: str,
    *,
    safe_reply_fn,
    apply_title_fn,
    base_result_fn,
) -> dict | None:
    """Gère export Excel depuis la conversation. Retourne None si hors contexte export."""
    from app.models.chat_message import ChatMessage, MessageSender

    last_bot = _last_bot_message(db, session.id)
    is_followup = last_bot is not None and last_bot.source == "export_prompt"
    if not is_followup and not is_export_intent(message):
        return None

    targets, reply, should_download = resolve_export_targets(db, session.id, user.id, message)

    if should_download and targets:
        count = len(targets)
        tn_label = targets[0] if count == 1 else f"{count} colis"
        download_reply = (
            f"J'exporte **{tn_label}** de cette conversation au format Excel.\n\n"
            "Le fichier contient uniquement les **données de suivi** (statut, localisation, ETA) — "
            "pas le texte de la conversation.\n\n"
            "Le téléchargement démarre automatiquement dans l'application."
        )
        bot = ChatMessage(
            session_id=session.id,
            sender=MessageSender.bot.value,
            source="export",
            message_text=safe_reply_fn(download_reply),
        )
        db.add(bot)
        apply_title_fn(session, message, bot_reply=download_reply, intent="export")
        db.commit()
        db.refresh(bot)
        result = base_result_fn(
            download_reply,
            session,
            source="export",
            intent="export",
            tracking_number=targets[0] if len(targets) == 1 else None,
            llm_provider=None,
        )
        result["export_download"] = {
            "session_id": session.id,
            "tracking_numbers": targets,
            "preset": "tracking",
            "include_events": True,
        }
        return result

    if reply:
        bot = ChatMessage(
            session_id=session.id,
            sender=MessageSender.bot.value,
            source="export_prompt",
            message_text=safe_reply_fn(reply),
        )
        db.add(bot)
        apply_title_fn(session, message, bot_reply=reply, intent="export")
        db.commit()
        db.refresh(bot)
        return base_result_fn(
            reply,
            session,
            source="export_prompt",
            intent="export",
            tracking_number=None,
            llm_provider=None,
        )

    return None


def _events_from_cache(raw_json: str | None) -> list[dict[str, Any]]:
    try:
        data = json.loads(raw_json or "{}")
    except json.JSONDecodeError:
        return []
    events = data.get("events")
    if isinstance(events, list):
        return [ev for ev in events if isinstance(ev, dict)]
    return []


def tracking_rows_to_shipment_payloads(
    db: Session,
    rows: list[TrackingRequest],
    *,
    include_events: bool = True,
) -> list[dict[str, Any]]:
    """Construit les données FedEx pour un export PDF client."""
    if not rows:
        return []
    tns = [r.tracking_number for r in rows if r.tracking_number]
    caches = {
        c.tracking_number: c
        for c in db.scalars(select(ShipmentCache).where(ShipmentCache.tracking_number.in_(tns))).all()
    }
    payloads: list[dict[str, Any]] = []
    for row in rows:
        tn = row.tracking_number
        cache = caches.get(tn)
        cached: dict[str, Any] = {}
        if cache and cache.raw_response_json:
            try:
                cached = json.loads(cache.raw_response_json)
            except json.JSONDecodeError:
                cached = {}
        events = _events_from_cache(cache.raw_response_json if cache else None) if include_events else []
        payloads.append(
            {
                "tracking_number": tn,
                "status": row.status or cached.get("status"),
                "current_location": row.current_location or cached.get("current_location"),
                "estimated_delivery": row.estimated_delivery or cached.get("estimated_delivery"),
                "actual_delivery": cached.get("actual_delivery"),
                "recipient": cached.get("recipient"),
                "origin_location": cached.get("origin_location"),
                "destination_location": cached.get("destination_location"),
                "events": events or list(cached.get("events") or []),
            }
        )
    return payloads


def generate_tracking_history_pdf_bytes(
    db: Session,
    *,
    user_id: int,
    session_id: int | None = None,
    tracking_numbers: list[str] | None = None,
    limit: int = 100,
    preset: str | None = None,
    include_events: bool = True,
) -> tuple[bytes, str]:
    """PDF récapitulatif des suivis (session, compte ou journée)."""
    from fpdf import FPDF

    from app.services.shipment_pdf_service import (
        _pdf_safe_text,
        generate_multi_shipment_history_pdf,
        generate_shipment_history_pdf,
    )

    rows = latest_tracking_rows(
        db,
        user_id=user_id,
        session_id=session_id,
        tracking_numbers=tracking_numbers,
        limit=limit,
    )
    if preset == "tracking_summary":
        today = datetime.now(timezone.utc).date()
        rows = [r for r in rows if r.created_at and r.created_at.date() == today]
        filename = "rapport-suivis-jour.pdf"
        title = "Rapport des suivis du jour"
    else:
        filename = "historique-suivi.pdf"
        title = "Historique de suivi FedEx"

    payloads = tracking_rows_to_shipment_payloads(db, rows, include_events=include_events)
    if payloads:
        if len(payloads) == 1:
            return generate_shipment_history_pdf(payloads[0]), filename
        return generate_multi_shipment_history_pdf(payloads), filename

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, _pdf_safe_text(title), ln=True)
    pdf.ln(4)
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(
        0,
        6,
        _pdf_safe_text(
            "Aucun suivi enregistré pour cette période. Consultez d'abord un numéro de colis dans le chat."
        ),
    )
    return pdf.output(), filename
