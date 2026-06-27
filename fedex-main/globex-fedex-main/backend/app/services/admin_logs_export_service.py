"""Export PDF / Excel des journaux d'activité admin (copilot admin)."""

from __future__ import annotations

import re
from datetime import datetime, timedelta, timezone
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from app.models.activity_log import ActivityLog


def parse_log_period_hours(task: str, *, default: int = 24) -> int:
    """Extrait une fenêtre temporelle depuis la consigne admin (ex. « 2h », « 5 jours », « 5jr »)."""
    t = (task or "").lower()
    m = re.search(r"(\d+)\s*(?:h|heures?)\b", t)
    if m:
        return min(max(int(m.group(1)), 1), 168)
    m = re.search(r"(\d+)\s*(?:jours?|jour|jr|j)\b", t) or re.search(r"(\d+)jr\b", t)
    if m:
        return min(int(m.group(1)) * 24, 168)
    if any(k in t for k in ("24h", "24 h", "journee", "journée", "aujourd", "today")):
        return 24
    if any(k in t for k in ("1h", "1 h", "une heure")):
        return 1
    return default


def _pdf_safe_text(value: str | None, *, max_len: int = 500) -> str:
    if not value:
        return "-"
    text = str(value).replace("\r", " ").replace("\n", " ").strip()
    text = text.replace("—", "-").replace("–", "-").replace("'", "'")
    try:
        text.encode("latin-1")
        return text[:max_len]
    except UnicodeEncodeError:
        return text.encode("latin-1", errors="replace").decode("latin-1")[:max_len]


def _pdf_write_line(pdf: Any, text: str, *, h: float = 7, bold: bool = False, size: int = 10) -> None:
    style = "B" if bold else ""
    pdf.set_font("Helvetica", style, size)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(pdf.epw, h, _pdf_safe_text(text))


def fetch_activity_logs(db: Session, *, hours: int, limit: int = 500) -> list[ActivityLog]:
    since = datetime.now(timezone.utc) - timedelta(hours=hours)
    return list(
        db.scalars(
            select(ActivityLog)
            .options(joinedload(ActivityLog.user), joinedload(ActivityLog.actor))
            .where(ActivityLog.created_at >= since)
            .order_by(ActivityLog.created_at.desc())
            .limit(limit)
        ).all()
    )


def generate_activity_logs_pdf(logs: list[ActivityLog], *, hours: int) -> tuple[bytes, str]:
    """Génère un PDF des logs d'activité. Retourne (bytes, filename)."""
    try:
        from fpdf import FPDF
    except ImportError as exc:
        raise RuntimeError("fpdf2 requis pour l'export PDF des logs.") from exc

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.add_page()
    _pdf_write_line(pdf, "Journaux d'activite Globex - Admin", h=10, bold=True, size=15)
    pdf.set_text_color(90, 90, 90)
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    _pdf_write_line(pdf, f"Periode : dernieres {hours}h - genere le {generated}")
    _pdf_write_line(pdf, f"Entrees : {len(logs)}")
    pdf.set_text_color(0, 0, 0)
    pdf.ln(4)

    if not logs:
        pdf.set_font("Helvetica", "I", 10)
        _pdf_write_line(pdf, "Aucune entree d'activite sur cette periode.")
    else:
        row_h = 7
        for row in logs:
            if pdf.get_y() > 265:
                pdf.add_page()

            ts = row.created_at.strftime("%d/%m %H:%M") if row.created_at else "-"
            actor = ""
            if row.actor and row.actor.email:
                actor = row.actor.email
            elif row.user and row.user.email:
                actor = row.user.email
            message = _pdf_safe_text(row.message, max_len=400)
            if actor:
                message = f"[{actor}] {message}"

            line = (
                f"{_pdf_safe_text(ts, max_len=16)} | "
                f"{_pdf_safe_text(row.level, max_len=10)} | "
                f"{_pdf_safe_text(row.action, max_len=24)} | "
                f"{message}"
            )
            _pdf_write_line(pdf, line, h=row_h, size=8)

    filename = f"activity-logs-{hours}h.pdf"
    out = pdf.output()
    if isinstance(out, bytearray):
        return bytes(out), filename
    if isinstance(out, bytes):
        return out, filename
    return out.encode("latin-1"), filename


def generate_activity_logs_excel(logs: list[ActivityLog], *, hours: int) -> tuple[bytes, str]:
    """Génère un Excel (date, titre, type de log). Retourne (bytes, filename)."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Journaux"
    headers = ["Date", "Titre", "Type de log"]
    ws.append(headers)
    header_fill = PatternFill(fill_type="solid", fgColor="4D148C")
    for i in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in logs:
        ts = row.created_at.strftime("%Y-%m-%d %H:%M:%S") if row.created_at else ""
        ws.append([ts, row.action or "", row.level or "INFO"])

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    for i, width in enumerate((20, 36, 14), start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    from io import BytesIO

    buf = BytesIO()
    wb.save(buf)
    filename = f"activity-logs-{hours}h.xlsx"
    return buf.getvalue(), filename


def build_logs_export_download(hours: int, *, fmt: str = "pdf") -> dict[str, Any]:
    if fmt == "xlsx":
        return {
            "preset": "admin_logs",
            "hours": hours,
            "filename": f"activity-logs-{hours}h.xlsx",
            "format": "xlsx",
        }
    return {
        "preset": "admin_logs",
        "hours": hours,
        "filename": f"activity-logs-{hours}h.pdf",
        "format": "pdf",
    }
