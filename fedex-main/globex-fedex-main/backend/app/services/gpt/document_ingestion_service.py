"""Ingestion de documents PDF/DOCX/TXT dans la base de connaissances GPT."""

from __future__ import annotations

import logging
import re
import uuid
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.models.gpt_definition import GptDefinition, KnowledgeChunk, KnowledgeCollection, KnowledgeDocument
from app.services.gpt.embedding_service import EmbeddingError, embed_texts
from app.services.gpt.text_chunking import split_text
from app.services.gpt.vector_store import store_chunk_embedding

logger = logging.getLogger(__name__)

_ALLOWED_EXT = frozenset({
    ".pdf", ".docx", ".txt", ".md", ".csv",
    ".xlsx", ".xls",
    ".jpg", ".jpeg", ".png", ".webp", ".gif",
})
_IMAGE_EXT = frozenset({".jpg", ".jpeg", ".png", ".webp", ".gif"})


class DocumentIngestError(Exception):
    pass


def _upload_root() -> Path:
    settings = get_settings()
    root = Path(__file__).resolve().parents[2] / settings.knowledge_upload_dir
    root.mkdir(parents=True, exist_ok=True)
    return root


def _extract_xlsx(data: bytes) -> str:
    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"## Feuille: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_image_text(data: bytes, *, mime_type: str) -> str:
    import base64

    from app.services.llm.providers import _gemini_generate

    b64 = base64.b64encode(data).decode("ascii")
    prompt = (
        "Transcris et décris le contenu de cette image pour indexation dans une base documentaire. "
        "Inclus tout texte visible, tableaux, chiffres et libellés. Réponds en français."
    )
    return _gemini_generate(
        prompt,
        max_output_tokens=2048,
        ui_language="fr",
        image_base64=b64,
        image_mime_type=mime_type,
    )


def extract_text_from_bytes(data: bytes, *, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        from io import BytesIO

        reader = PdfReader(BytesIO(data))
        parts: list[str] = []
        for page in reader.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text.strip())
        return "\n\n".join(parts)

    if ext == ".docx":
        from io import BytesIO

        from docx import Document

        doc = Document(BytesIO(data))
        return "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())

    if ext in {".xlsx", ".xls"}:
        return _extract_xlsx(data)

    if ext in _IMAGE_EXT:
        mime = _mime_for_ext(ext)
        return _extract_image_text(data, mime_type=mime)

    if ext in {".txt", ".md", ".csv"}:
        for encoding in ("utf-8", "latin-1", "cp1252"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", errors="replace")

    raise DocumentIngestError(f"Format non supporté : {ext or 'inconnu'}")


def _slugify_filename(name: str) -> str:
    base = Path(name).stem.lower()
    base = re.sub(r"[^a-z0-9]+", "-", base).strip("-")
    return base[:48] or "document"


def get_or_create_upload_collection(
    db: Session,
    *,
    gpt: GptDefinition,
    language: str,
    title: str | None = None,
) -> KnowledgeCollection:
    slug = f"uploads-{language}"
    row = db.scalar(
        select(KnowledgeCollection).where(
            KnowledgeCollection.gpt_id == gpt.id,
            KnowledgeCollection.slug == slug,
        )
    )
    if row is None:
        row = KnowledgeCollection(
            gpt_id=gpt.id,
            slug=slug,
            title=title or f"Documents uploadés ({language})",
            language=language,
            source_type="upload",
            description="Documents PDF/DOCX indexés par l'admin",
        )
        db.add(row)
        db.flush()
        _attach_collection_to_gpt(gpt, row.id)
    return row


def _attach_collection_to_gpt(gpt: GptDefinition, collection_id: int) -> None:
    import json

    try:
        ids = json.loads(gpt.knowledge_collection_ids_json or "[]")
        if not isinstance(ids, list):
            ids = []
    except json.JSONDecodeError:
        ids = []
    if collection_id not in ids:
        ids.append(collection_id)
        gpt.knowledge_collection_ids_json = json.dumps(ids)


def ingest_uploaded_document(
    db: Session,
    *,
    gpt: GptDefinition,
    file_bytes: bytes,
    filename: str,
    language: str = "fr",
    uploaded_by_user_id: int | None = None,
) -> KnowledgeDocument:
    """Extrait, découpe, embed et indexe un document uploadé."""
    settings = get_settings()
    ext = Path(filename).suffix.lower()
    if ext not in _ALLOWED_EXT:
        raise DocumentIngestError(f"Extension autorisée : {', '.join(sorted(_ALLOWED_EXT))}")

    max_bytes = settings.knowledge_upload_max_mb * 1024 * 1024
    if len(file_bytes) > max_bytes:
        raise DocumentIngestError(f"Fichier trop volumineux (max {settings.knowledge_upload_max_mb} Mo).")

    collection = get_or_create_upload_collection(db, gpt=gpt, language=language)
    stored_name = f"{uuid.uuid4().hex}_{Path(filename).name}"
    stored_path = _upload_root() / stored_name
    stored_path.write_bytes(file_bytes)

    doc = KnowledgeDocument(
        collection_id=collection.id,
        original_filename=filename,
        stored_path=str(stored_path),
        mime_type=_mime_for_ext(ext),
        file_size_bytes=len(file_bytes),
        status="processing",
        uploaded_by_user_id=uploaded_by_user_id,
    )
    db.add(doc)
    db.flush()

    try:
        text_content = extract_text_from_bytes(file_bytes, filename=filename)
        if not text_content.strip():
            raise DocumentIngestError("Aucun texte extractible du document.")

        pieces = split_text(
            text_content,
            chunk_size=settings.gpt_chunk_size,
            overlap=settings.gpt_chunk_overlap,
        )
        if not pieces:
            raise DocumentIngestError("Découpage vide après extraction.")

        embeddings: list[list[float]] = []
        embedding_warning = ""
        try:
            embeddings = embed_texts(pieces, task_type="RETRIEVAL_DOCUMENT")
        except EmbeddingError as exc:
            embedding_warning = str(exc)[:240]
            logger.warning("Indexation sans vecteurs pour %s : %s", filename, exc)
            embeddings = [[] for _ in pieces]

        slug_base = _slugify_filename(filename)

        for idx, piece in enumerate(pieces):
            vector = embeddings[idx] if idx < len(embeddings) else []
            chunk_key = f"{slug_base}-{doc.id}-{idx}"
            title = f"{Path(filename).stem} — partie {idx + 1}"
            chunk = KnowledgeChunk(
                collection_id=collection.id,
                document_id=doc.id,
                chunk_index=idx,
                chunk_key=chunk_key,
                title=title,
                content=piece,
                keywords=_auto_keywords(piece),
                source_ref=filename,
                language=language,
            )
            db.add(chunk)
            db.flush()
            if vector:
                store_chunk_embedding(db, chunk_id=chunk.id, vector=vector)

        doc.status = "indexed"
        doc.chunk_count = len(pieces)
        doc.error_message = embedding_warning
        db.commit()
        db.refresh(doc)
        return doc
    except DocumentIngestError as exc:
        db.rollback()
        doc = db.get(KnowledgeDocument, doc.id)
        if doc is not None:
            doc.status = "failed"
            doc.error_message = str(exc)[:500]
            db.commit()
        raise
    except Exception as exc:
        db.rollback()
        doc = db.get(KnowledgeDocument, doc.id)
        if doc is not None:
            doc.status = "failed"
            doc.error_message = str(exc)[:500]
            db.commit()
        logger.exception("Ingestion document %s", filename)
        raise DocumentIngestError(f"Indexation échouée : {exc}") from exc


def _mime_for_ext(ext: str) -> str:
    return {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".xls": "application/vnd.ms-excel",
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".csv": "text/csv",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".webp": "image/webp",
        ".gif": "image/gif",
    }.get(ext, "application/octet-stream")


def _auto_keywords(text: str) -> str:
    words = re.findall(r"[a-zA-Zàâäéèêëïîôùûüç0-9]{4,}", text.lower())
    unique = list(dict.fromkeys(words))[:12]
    return ", ".join(unique)


def reindex_chunks_without_embeddings(db: Session, *, gpt_id: int | None = None, limit: int = 200) -> int:
    """Recalcule les embeddings manquants (seed JSON + migrations)."""
    from app.models.gpt_definition import KnowledgeCollection

    stmt = (
        select(KnowledgeChunk)
        .join(KnowledgeCollection)
        .where(
            (KnowledgeChunk.embedding_json.is_(None)) | (KnowledgeChunk.embedding_json == "")
        )
        .limit(limit)
    )
    if gpt_id is not None:
        stmt = stmt.where(KnowledgeCollection.gpt_id == gpt_id)

    chunks = list(db.scalars(stmt).all())
    if not chunks:
        return 0

    texts = [f"{c.title}\n{c.content}" for c in chunks]
    vectors = embed_texts(texts, task_type="RETRIEVAL_DOCUMENT")
    for chunk, vector in zip(chunks, vectors):
        store_chunk_embedding(db, chunk_id=chunk.id, vector=vector)
    db.commit()
    return len(chunks)
