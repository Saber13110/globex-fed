"""Extraction texte depuis fichiers (partagé admin GPT + client Phase 9)."""

from __future__ import annotations

import base64
from io import BytesIO
from pathlib import Path

from app.services.llm.providers import _gemini_generate, gemini_api_key_usable

_IMAGE_EXT = frozenset({".jpg", ".jpeg", ".png", ".webp", ".gif"})
_CLIENT_EXT = frozenset({".pdf", ".xlsx", ".xls", *_IMAGE_EXT})
_ADMIN_EXTRA_EXT = frozenset({".docx", ".txt", ".md", ".csv"})

_OCR_PROMPT = (
    "Transcris et décris le contenu de cette image pour indexation dans une base documentaire. "
    "Inclus tout texte visible, tableaux, chiffres et libellés. Réponds en français."
)


class DocumentExtractError(Exception):
    pass


def mime_for_ext(ext: str) -> str:
    return {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".xls": "application/vnd.ms-excel",
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".csv": "text/csv",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".webp": "image/webp",
        ".gif": "image/gif",
    }.get(ext, "application/octet-stream")


def is_client_document_ext(ext: str) -> bool:
    return ext.lower() in _CLIENT_EXT


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"## Feuille: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_image_text(data: bytes, *, mime_type: str) -> str:
    if not gemini_api_key_usable():
        raise DocumentExtractError(
            "OCR indisponible — configurez GEMINI_API_KEY pour lire les images (comme l'upload admin)."
        )
    b64 = base64.b64encode(data).decode("ascii")
    return _gemini_generate(
        _OCR_PROMPT,
        max_output_tokens=2048,
        ui_language="fr",
        image_base64=b64,
        image_mime_type=mime_type,
    )


def extract_text_from_bytes(data: bytes, *, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        parts: list[str] = []
        for page in reader.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text.strip())
        return "\n\n".join(parts)

    if ext == ".docx":
        from docx import Document

        doc = Document(BytesIO(data))
        return "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())

    if ext in {".xlsx", ".xls"}:
        return _extract_xlsx(data)

    if ext in _IMAGE_EXT:
        return _extract_image_text(data, mime_type=mime_for_ext(ext))

    if ext in {".txt", ".md", ".csv"}:
        for encoding in ("utf-8", "latin-1", "cp1252"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", errors="replace")

    raise DocumentExtractError(f"Format non supporté : {ext or 'inconnu'}")
